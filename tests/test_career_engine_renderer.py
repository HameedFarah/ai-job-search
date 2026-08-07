from __future__ import annotations

import json
import shutil
import subprocess
from pathlib import Path

import pytest
from docx import Document

from career_engine.bundle import build_bundle
from career_engine.core import decide_route, match_evidence, normalize_job, score_fit
from career_engine.generation import create_generation_packet
from career_engine.pipeline import finalize_render, import_generated, prepare, status as pipeline_status
from career_engine.renderer import (
    _all_paragraphs,
    _libreoffice_binary,
    build_render_input,
    convert_docx_to_pdf,
    render_and_verify,
    render_docx,
    render_tooling,
    verify_pdf,
)
from tests.test_career_engine_v1 import engine_root, job_payload, valid_application  # noqa: F401

REPO = Path(__file__).resolve().parents[1]


def seven_bullet_application(packet: dict) -> dict:
    claims = [item["id"] for item in packet["selected_claims"]]
    metrics = list(dict.fromkeys(packet["selected_metric_claim_ids"] + claims))[:6]
    base = valid_application(packet)
    base["current_role_bullets"] = [
        {"text": f"Delivered {label} across complex Saudi programmes with measurable impact.", "claim_ids": [claims[i % len(claims)]]}
        for i, label in enumerate(
            ("design governance", "multidisciplinary coordination", "client leadership",
             "value engineering", "quality assurance", "project delivery", "team leadership")
        )
    ]
    base["metric_claim_ids"] = metrics
    return base


def test_all_paragraphs_finds_template_placeholders(engine_root: Path) -> None:
    template = engine_root / "projects/job-automation/config/career-engine.v1.json"
    template_path = REPO / json.loads(template.read_text())["template"]["repository_path"]
    document = Document(template_path)
    texts = [p.text for p in _all_paragraphs(document)]
    assert sum(1 for t in texts if t.startswith("• [ACHIEVEMENT")) == 7
    assert sum(1 for t in texts if t in {f"[M{i}]" for i in range(1, 7)}) == 6
    assert sum(1 for t in texts if t.startswith("[VACANCY-RELEVANT")) == 6
    assert sum(1 for t in texts if t.startswith("[EVIDENCE CARD")) == 4
    assert sum(1 for t in texts if t.startswith("[One concise evidence statement")) == 4
    assert sum(1 for t in texts if t.startswith("Representative KSA context")) == 1
    assert any("[TARGET ROLE HEADLINE]" in t for t in texts)
    assert any("[TAILORED PROFILE" in t for t in texts)


def test_build_render_input_writes_render_input_json(job_payload: dict, engine_root: Path) -> None:
    bundle = build_bundle(engine_root)
    normalized = normalize_job(job_payload, bundle["taxonomy"])
    matches = match_evidence(normalized, bundle)
    score = score_fit(normalized, matches, bundle)
    route = decide_route(normalized, bundle)
    packet = create_generation_packet(job_id="render-test-job", normalized_job=normalized, matches=matches, score=score, route=route, bundle=bundle)
    application = seven_bullet_application(packet)
    result = build_render_input("render-test-job", application, packet, root=engine_root)
    path = Path(result["path"])
    assert path.is_file()
    data = json.loads(path.read_text())
    assert data["job_id"] == "render-test-job"
    assert data["bundle_hash"] == packet["bundle_hash"]
    assert data["layout"]["page_limit"] == 2
    assert data["layout"]["headshot_required"] is True
    assert data["outward_filename"] == packet["outward_filename"]


def test_render_docx_replaces_placeholders(job_payload: dict, engine_root: Path) -> None:
    bundle = build_bundle(engine_root)
    normalized = normalize_job(job_payload, bundle["taxonomy"])
    matches = match_evidence(normalized, bundle)
    score = score_fit(normalized, matches, bundle)
    route = decide_route(normalized, bundle)
    packet = create_generation_packet(job_id="render-test-job", normalized_job=normalized, matches=matches, score=score, route=route, bundle=bundle)
    application = seven_bullet_application(packet)
    result = render_docx("render-test-job", application, packet, root=engine_root)
    assert Path(result["docx"]).is_file()
    assert result["sha256"]
    assert result["outward_filename"] == Path(packet["outward_filename"]).with_suffix(".docx").name
    rendered = Document(result["docx"])
    texts = [p.text for p in _all_paragraphs(rendered)]
    assert application["headline"] in texts
    rendered_text = "\n".join(texts)
    assert "• [ACHIEVEMENT" not in rendered_text
    assert "[M1]" not in rendered_text
    assert "[TARGET ROLE HEADLINE]" not in rendered_text
    assert application["earlier_role_bullets"][0]["text"] in rendered_text
    assert application["earlier_role_bullets"][-1]["text"] in rendered_text
    assert "• Directed end-to-end design and project delivery, from client briefing" not in rendered_text
    assert "• Led multidisciplinary teams, consultants, contractors and stakeholders, aligning technical quality" not in rendered_text


def test_finalize_render_moves_tracker_to_owner_approval(
    job_payload: dict,
    engine_root: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    payload = dict(job_payload)
    payload.update({
        "live_status": "live",
        "live_verified_at": "2026-08-03T10:00:00+00:00",
        "live_verification_source": "official employer careers page",
    })
    state = prepare(payload, root=engine_root, actor="chatgpt")
    packet_path = Path(state["outputs"]["generation_packet"])
    packet = json.loads(packet_path.read_text(encoding="utf-8"))
    application = valid_application(packet)
    pending = packet_path.parent / "generated_application.pending.json"
    pending.write_text(json.dumps(application, ensure_ascii=False, indent=2), encoding="utf-8")
    imported = import_generated(state["job_id"], pending, root=engine_root, actor="chatgpt")
    assert imported["valid"] is True

    final_docx = packet_path.parent / "Abdelhamid_Farah_CV_Test.docx"
    final_pdf = packet_path.parent / "Abdelhamid_Farah_CV_Test.pdf"
    ats_docx = packet_path.parent / "Abdelhamid_Farah_CV_Test_ATS.docx"
    ats_pdf = packet_path.parent / "Abdelhamid_Farah_CV_Test_ATS.pdf"
    final_docx.write_bytes(b"docx")
    final_pdf.write_bytes(b"pdf")
    ats_docx.write_bytes(b"ats-docx")
    ats_pdf.write_bytes(b"ats-pdf")
    monkeypatch.setattr(
        "career_engine.pipeline.render_and_verify",
        lambda *args, **kwargs: {
            "valid": True,
            "docx": {"docx": str(final_docx), "sha256": "docx-sha"},
            "verification": {"pdf": str(final_pdf), "sha256": "pdf-sha", "valid": True},
        },
    )
    monkeypatch.setattr(
        "career_engine.pipeline.render_ats_and_verify",
        lambda *args, **kwargs: {
            "valid": True,
            "docx": {"docx": str(ats_docx), "sha256": "ats-docx-sha"},
            "verification": {"pdf": str(ats_pdf), "sha256": "ats-pdf-sha", "valid": True},
        },
    )

    result = finalize_render(state["job_id"], root=engine_root, actor="chatgpt")
    assert result["valid"] is True
    assert result["submission_package"]["default_resume_variant"] == "ats-linear"
    assert result["submission_package"]["selected_resume_variant"] == "ats-linear"
    assert result["submission_package"]["owner_override"] is False
    assert result["submission_package"]["attachment_count"] == 0
    assert result["submission_package"]["email_account"] == "hameedo@gmail.com"
    assert result["submission_package"]["email_sender"] == "hameedfarah@gmail.com"
    tracker_state = pipeline_status(state["job_id"], root=engine_root)
    assert tracker_state["processing_state"]["status"] == "awaiting_owner_approval"
    assert tracker_state["processing_state"]["owner"] == "owner"
    assert tracker_state["processing_state"]["external_action_allowed"] is False
    assert tracker_state["processing_state"]["selected_resume_variant"] == "ats-linear"
    assert tracker_state["processing_state"]["submission_package"]["selected_resume_variant"] == "ats-linear"
    record = tracker_state["job"]
    assert record["processing_status"] == "awaiting_owner_approval"
    assert record["owner"] == "owner"
    record_path = engine_root / "projects/job-automation/data/jobs" / f"{state['job_id']}.json"
    saved_record = json.loads(record_path.read_text(encoding="utf-8"))
    assert saved_record["submission_package"]["selected_resume_variant"] == "ats-linear"


def test_finalize_render_persisted_preview_override_selects_single_cv(
    job_payload: dict,
    engine_root: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """A persisted per-job preview override changes the single selected CV variant.

    The portal route defaults to ATS Linear; the persisted override selects the
    Modern Executive Sidebar variant while both variants remain generated.
    """
    payload = dict(job_payload)
    payload.update({
        "live_status": "live",
        "live_verified_at": "2026-08-03T10:00:00+00:00",
        "live_verification_source": "official employer careers page",
    })
    state = prepare(payload, root=engine_root, actor="chatgpt")
    packet_path = Path(state["outputs"]["generation_packet"])
    packet = json.loads(packet_path.read_text(encoding="utf-8"))
    application = valid_application(packet)
    pending = packet_path.parent / "generated_application.pending.json"
    pending.write_text(json.dumps(application, ensure_ascii=False, indent=2), encoding="utf-8")
    assert import_generated(state["job_id"], pending, root=engine_root, actor="chatgpt")["valid"] is True

    from career_engine.config import load_config
    from career_engine.pipeline import _load_tracker

    _, paths = load_config(engine_root)
    tracker = _load_tracker(paths)
    tracker.update_job(
        state["job_id"],
        {"resume_template_override": "modern-executive-sidebar"},
        comment="Owner preview override selects the Modern Executive Sidebar CV",
        actor="owner",
    )

    final_docx = packet_path.parent / "Abdelhamid_Farah_CV_Test.docx"
    final_pdf = packet_path.parent / "Abdelhamid_Farah_CV_Test.pdf"
    ats_docx = packet_path.parent / "Abdelhamid_Farah_CV_Test_ATS.docx"
    ats_pdf = packet_path.parent / "Abdelhamid_Farah_CV_Test_ATS.pdf"
    for path in (final_docx, final_pdf, ats_docx, ats_pdf):
        path.write_bytes(b"content")
    monkeypatch.setattr(
        "career_engine.pipeline.render_and_verify",
        lambda *args, **kwargs: {
            "valid": True,
            "docx": {"docx": str(final_docx), "sha256": "docx-sha"},
            "verification": {"pdf": str(final_pdf), "sha256": "pdf-sha", "valid": True},
        },
    )
    monkeypatch.setattr(
        "career_engine.pipeline.render_ats_and_verify",
        lambda *args, **kwargs: {
            "valid": True,
            "docx": {"docx": str(ats_docx), "sha256": "ats-docx-sha"},
            "verification": {"pdf": str(ats_pdf), "sha256": "ats-pdf-sha", "valid": True},
        },
    )

    result = finalize_render(state["job_id"], root=engine_root, actor="chatgpt")
    assert result["valid"] is True
    package = result["submission_package"]
    assert package["default_resume_variant"] == "ats-linear"
    assert package["selected_resume_variant"] == "modern-executive-sidebar"
    assert package["owner_override"] is True
    assert package["selected_cv_pdf"].endswith("Abdelhamid_Farah_CV_Test.pdf")
    assert package["selected_cv_docx"].endswith("Abdelhamid_Farah_CV_Test.docx")
    # Both variants remain generated and visible in the dashboard artifacts.
    record_path = engine_root / "projects/job-automation/data/jobs" / f"{state['job_id']}.json"
    record = json.loads(record_path.read_text(encoding="utf-8"))
    generated = record["generated_artifacts"]
    assert any(item["type"] == "ats_pdf" for item in generated)
    assert any(item["type"] == "final_pdf" for item in generated)
    assert record["processing_state"]["selected_resume_variant"] == "modern-executive-sidebar"
    assert record["submission_package"]["selected_cv_pdf"].endswith("Abdelhamid_Farah_CV_Test.pdf")


def test_render_tooling_accepts_explicit_binary_override(monkeypatch: pytest.MonkeyPatch) -> None:
    resolved = shutil.which("true")
    assert resolved
    binary = Path(resolved)
    monkeypatch.setenv("CAREER_ENGINE_LIBREOFFICE", str(binary))
    tooling = render_tooling()
    assert tooling["libreoffice"] is True
    assert tooling["libreoffice_path"] == str(binary.resolve())
    assert tooling["pdf_conversion_available"] is True


def _fake_soffice(path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("#!/bin/sh\n")
    path.chmod(0o755)
    return path


def test_libreoffice_binary_discovers_newest_user_local_soffice(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("CAREER_ENGINE_LIBREOFFICE", raising=False)
    monkeypatch.setenv("PATH", str(tmp_path / "no-bin"))
    home = tmp_path / "home"
    older = _fake_soffice(home / ".local/opt/libreoffice-26.2.4/opt/libreoffice26.2/program/soffice")
    newer = _fake_soffice(home / ".local/opt/libreoffice-26.2.5/opt/libreoffice26.2/program/soffice")
    assert _libreoffice_binary(home=home) == str(newer.resolve())
    assert _libreoffice_binary(home=home) != str(older.resolve())


def test_libreoffice_binary_prefers_system_soffice_over_user_local(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("CAREER_ENGINE_LIBREOFFICE", raising=False)
    system_bin = tmp_path / "system-bin"
    system_bin.mkdir()
    system = _fake_soffice(system_bin / "soffice")
    monkeypatch.setenv("PATH", str(system_bin))
    home = tmp_path / "home"
    user_local = _fake_soffice(home / ".local/opt/libreoffice-26.2.5/opt/libreoffice26.2/program/soffice")
    assert _libreoffice_binary(home=home) == str(system.resolve())
    assert _libreoffice_binary(home=home) != str(user_local.resolve())


def test_libreoffice_binary_env_override_wins_over_system_and_user_local(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    override = _fake_soffice(tmp_path / "override" / "soffice")
    monkeypatch.setenv("CAREER_ENGINE_LIBREOFFICE", str(override))
    system_bin = tmp_path / "system-bin"
    system_bin.mkdir()
    system = _fake_soffice(system_bin / "soffice")
    monkeypatch.setenv("PATH", str(system_bin))
    home = tmp_path / "home"
    _fake_soffice(home / ".local/opt/libreoffice-26.2.5/opt/libreoffice26.2/program/soffice")
    assert _libreoffice_binary(home=home) == str(override.resolve())
    assert _libreoffice_binary(home=home) != str(system.resolve())


def test_libreoffice_binary_returns_empty_when_unavailable(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.delenv("CAREER_ENGINE_LIBREOFFICE", raising=False)
    monkeypatch.setenv("PATH", str(tmp_path / "no-bin"))
    assert _libreoffice_binary(home=tmp_path / "empty-home") == ""


def test_convert_docx_to_pdf_builds_headless_command(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    real_true = shutil.which("true")
    assert real_true
    monkeypatch.setenv("CAREER_ENGINE_LIBREOFFICE", real_true)
    profile_base = tmp_path / "profiles"

    def fake_profile_dir() -> Path:
        profile_base.mkdir(parents=True, exist_ok=True)
        return profile_base

    monkeypatch.setattr("career_engine.renderer._libreoffice_profile_dir", fake_profile_dir)
    docx_path = tmp_path / "cv.docx"
    docx_path.write_bytes(b"fake docx")
    output_dir = tmp_path / "out"
    pdf = output_dir / "cv.pdf"
    pdf.parent.mkdir(parents=True)
    pdf.write_bytes(b"fake pdf")

    captured: dict = {}

    def fake_run(command: list[str], **kwargs: object) -> subprocess.CompletedProcess:
        captured["command"] = list(command)
        captured["env"] = kwargs.get("env")
        return subprocess.CompletedProcess(command, 0, stdout="convert ok", stderr="")

    monkeypatch.setattr("career_engine.renderer.subprocess.run", fake_run)
    result = convert_docx_to_pdf(docx_path, output_dir)
    assert result["converted"] is True
    assert result["returncode"] == 0
    assert result["pdf"] == str(pdf)
    command = captured["command"]
    assert command[0] == str(Path(real_true).resolve())
    assert "--headless" in command
    assert "--norestore" in command
    install_flag = next(item for item in command if item.startswith("-env:UserInstallation="))
    profile_uri = install_flag.split("=", 1)[1]
    assert profile_uri.startswith("file://")
    assert "career-engine-lo-" in profile_uri
    assert captured["env"]["TMPDIR"] == profile_uri.removeprefix("file://")
    assert command[command.index("--convert-to") + 1] == "pdf"
    assert command[command.index("--outdir") + 1] == str(output_dir)
    assert command[-1] == str(docx_path)


def test_doctor_reports_libreoffice_tooling(engine_root: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from career_engine.cli import doctor

    real_true = shutil.which("true")
    assert real_true
    monkeypatch.setenv("CAREER_ENGINE_LIBREOFFICE", real_true)
    build_bundle(engine_root)
    result = doctor(engine_root)
    tooling = result["render_tooling"]
    assert tooling["libreoffice"] is True
    assert tooling["libreoffice_path"] == str(Path(real_true).resolve())
    assert tooling["pdf_conversion_available"] is True


def test_render_and_verify_requires_libreoffice_for_pdf(job_payload: dict, engine_root: Path) -> None:
    bundle = build_bundle(engine_root)
    normalized = normalize_job(job_payload, bundle["taxonomy"])
    matches = match_evidence(normalized, bundle)
    score = score_fit(normalized, matches, bundle)
    route = decide_route(normalized, bundle)
    packet = create_generation_packet(job_id="render-test-job", normalized_job=normalized, matches=matches, score=score, route=route, bundle=bundle)
    application = seven_bullet_application(packet)
    result = render_and_verify("render-test-job", application, packet, root=engine_root)
    assert result["docx"]["docx"]
    if not render_tooling()["pdf_conversion_available"]:
        assert result["valid"] is False
        assert result["conversion"]["converted"] is False
        assert "libreoffice_missing" in str(result["conversion"].get("blocker", ""))
    else:
        assert result["valid"] is True
        assert result["verification"]["page_count"] == 2


@pytest.mark.skipif(not shutil.which("pdfinfo") or not shutil.which("pdftotext"), reason="poppler tools unavailable")
def test_verify_pdf_accepts_artifact_pdf(engine_root: Path) -> None:
    artifact = REPO / "projects/job-automation/artifacts/5a531dd6cfca13213694/Abdelhamid_Farah_CV_Design_Governance_Manager.pdf"
    if not artifact.is_file():
        pytest.skip("artifact PDF not present")
    # Artifacts are untracked runtime data. A locally present artifact produced
    # under the superseded outward-email policy cannot satisfy the current
    # required-identity check, so skip it instead of failing on stale data.
    extracted = subprocess.run(
        ["pdftotext", str(artifact), "-"], capture_output=True, text=True, check=False
    ).stdout
    if "hameedfarah@gmail.com" not in extracted:
        pytest.skip("artifact PDF predates the current hameedfarah@gmail.com outward-email policy")
    result = verify_pdf(artifact, root=engine_root)
    assert result["valid"] is True
    assert result["page_count"] == 2
    assert result["text_characters"] > 1000


@pytest.mark.skipif(not shutil.which("pdfinfo") or not shutil.which("pdftotext"), reason="poppler tools unavailable")
def test_verify_pdf_rejects_page_count_mismatch(engine_root: Path) -> None:
    artifact = REPO / "projects/job-automation/artifacts/5a531dd6cfca13213694/Abdelhamid_Farah_Qiddiya_Manager_Design_Governance_Cover_Letter.pdf"
    if not artifact.is_file():
        pytest.skip("artifact PDF not present")
    result = verify_pdf(artifact, root=engine_root)
    assert result["valid"] is False
    assert any(item["code"] == "page_count" for item in result["findings"])
